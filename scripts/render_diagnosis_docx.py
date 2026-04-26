#!/usr/bin/env python3

import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def parse_args():
    parser = argparse.ArgumentParser(
        description="Render a diagnosis report JSON file into a DOCX report."
    )
    parser.add_argument("--input", required=True, help="Path to diagnosis report JSON.")
    parser.add_argument("--output", required=True, help="Path to output DOCX file.")
    return parser.parse_args()


def load_json(path: Path):
    return json.loads(path.read_text())


def set_normal_style(document: Document):
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)


def add_bullets(document: Document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_diagnosis_items(document: Document, items):
    if not items:
        document.add_paragraph("当前版本暂无相关内容。")
        return

    for item in items:
        document.add_heading(item["title"], level=3)
        document.add_paragraph(f"结论类型：{item['finding_type']}")
        document.add_paragraph(f"依据归属：{item['basis_type']}")
        document.add_paragraph(f"严重度：{item['severity']}")
        document.add_paragraph(f"诊断主题：{item['category_l1']} / {item['category_l2']}")
        document.add_paragraph(f"摘要：{item['summary']}")
        if item.get("detailed_analysis"):
            document.add_paragraph(f"说明：{item['detailed_analysis']}")

        document.add_paragraph("证据位置：")
        for evidence in item["evidence_refs"]:
            region_id = evidence["region_id"] or "未分区"
            document.add_paragraph(
                f"第{evidence['page_index'] + 1}页 / {region_id} / {evidence['description']}",
                style="List Bullet",
            )

        document.add_paragraph(f"规则引用：{', '.join(item['rule_refs'])}")
        document.add_paragraph(f"建议：{item['recommendation']}")
        if item.get("confidence") is not None:
            document.add_paragraph(f"置信度：{item['confidence']:.2f}")


def add_evidence_images(document: Document, images):
    if not images:
        document.add_paragraph("当前版本暂无证据图。")
        return

    for image in images:
        document.add_heading(image["caption"], level=3)
        image_path = image.get("path")
        if image_path and Path(image_path).exists():
            document.add_picture(image_path, width=Inches(6.2))
        else:
            document.add_paragraph("证据图文件暂不可用。")


def main():
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    report = load_json(input_path)

    document = Document()
    set_normal_style(document)

    document.add_heading(f"{report['project_info']['project_name']} 诊断报告", level=1)

    document.add_heading("总体判断", level=2)
    document.add_paragraph(report["executive_summary"]["overall_view"])

    if report["executive_summary"].get("key_strengths"):
        document.add_heading("主要优势", level=3)
        add_bullets(document, report["executive_summary"]["key_strengths"])

    if report["executive_summary"].get("key_risks"):
        document.add_heading("主要风险", level=3)
        add_bullets(document, report["executive_summary"]["key_risks"])

    document.add_heading("错误指出意见", level=2)
    add_diagnosis_items(document, report["issues"])

    document.add_heading("优化性意见", level=2)
    add_diagnosis_items(document, report["optimizations"])

    if report.get("citations"):
        document.add_heading("依据索引", level=2)
        for citation in report["citations"]:
            document.add_paragraph(
                f"{citation['citation_id']} | {citation['source_type']} | {citation['source_file']} | {citation['source_locator']} | {citation['quoted_or_paraphrased_text']}",
                style="List Bullet",
            )

    appendix = report.get("appendix", {})
    document.add_heading("证据附录", level=2)
    add_evidence_images(document, appendix.get("evidence_images", []))

    document.save(str(output_path))
    print(json.dumps({"status": "ok", "output_path": str(output_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
